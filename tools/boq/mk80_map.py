# ملف مراجعة: مقابلة كل بند في نطاق العمل (SOW) ببند جدول كميات مركز فاطمة العفالق
import json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
sys.path.insert(0, os.path.dirname(__file__))
from mk80_fatima import L

SOW, OUT = sys.argv[1], sys.argv[2]
KEY = {x['k']: i + 1 for i, x in enumerate(L)}
M = {17: 'LS-41', 18: 'LS-33', 19: 'LS-33', 20: 'LS-33', 23: 'LS-40', 24: 'LS-40', 28: 'LS-44', 29: 'LS-44', 30: 'LS-45',
  33: 'LS-89', 35: 'LS-89', 36: 'LS-89', 46: 'LS-89', 39: 'LS-90', 40: 'LS-90', 41: 'LS-90', 48: 'LS-90', 44: 'LS-91', 45: 'LS-91', 47: 'LS-65',
  53: 'LS-46', 56: 'LS-46', 57: 'LS-46', 58: 'LS-46', 59: 'LS-46', 60: 'LS-47', 61: 'LS-47', 62: 'LS-47', 64: 'LS-47', 63: 'LS-48',
  68: 'LS-49', 69: 'LS-49', 73: 'LS-53', 74: 'LS-53', 75: 'LOUV', 76: 'LOUV', 85: 'LS-62',
  95: 'LS-46', 96: 'LS-48', 97: 'LS-50', 98: 'LS-51', 102: 'LS-54', 103: 'LS-54', 104: 'LS-56', 105: 'LS-56', 106: 'XRAY', 107: 'XRAY', 108: 'LS-55', 109: 'XRAY',
  110: 'LS-56', 111: 'LS-56', 112: 'LS-56', 113: 'LS-56', 118: 'LS-58', 121: 'LS-58', 122: 'LS-59', 123: 'LS-58', 124: 'LS-58', 125: 'LS-51', 126: 'LS-59',
  129: 'LS-60', 130: 'LS-60', 131: 'LS-61', 134: 'PART', 137: 'LS-57', 138: 'LS-57', 139: 'LS-59', 140: 'LS-61', 141: 'XRAY', 142: 'LS-58', 145: 'LS-59',
  150: 'LS-62', 153: 'LS-66', 154: 'LS-65', 157: 'LS-68', 158: 'LS-68', 161: 'LS-63', 165: None, 166: None, 170: 'LS-93', 174: 'LS-69', 175: 'LS-69',
  180: 'LS-70', 181: 'LS-70', 183: 'LS-70', 185: 'LS-70', 186: 'LS-70', 182: 'LS-72', 189: 'LS-72', 184: 'LS-71',
  194: 'LS-73', 195: 'LS-73', 196: 'LS-73', 201: 'LS-73', 204: 'LS-73', 205: 'LS-73', 208: 'LS-73', 197: 'LS-48', 198: 'LS-75', 199: 'LS-75', 200: 'LS-75', 213: 'LS-75',
  206: 'LS-74', 207: 'LS-74', 211: 'LS-91', 212: 'LS-91', 214: 'LS-91', 215: 'LS-91',
  255: 'LS-77', 256: 'LS-77', 258: 'LS-77', 259: 'LS-77', 260: 'LS-77', 261: 'LS-77', 262: 'LS-77', 263: 'LS-77', 257: 'LS-78', 264: 'LS-31', 268: 'BMS',
  272: 'MV', 273: 'MV', 281: 'MV', 274: 'GEN', 275: 'GEN', 276: 'LS-79', 277: 'LS-79', 278: 'LS-79', 280: 'LS-79', 279: 'IPS', 282: 'LS-81', 283: 'LS-81',
  284: 'LS-80', 285: 'LS-80', 286: 'LS-80', 287: 'LS-80', 288: 'LS-80', 289: 'LS-80', 290: 'LS-82', 295: 'LS-83', 296: 'LS-83', 297: 'LS-83', 300: 'MATV',
  303: 'LS-85', 304: 'LS-85', 307: 'CLOCK', 308: 'CLOCK', 311: 'AV', 312: 'AV', 315: 'LS-88', 320: 'ACS', 323: 'LS-87', 324: 'LS-87', 325: 'LS-87',
  328: 'LS-86', 329: 'LS-86', 330: 'LS-86', 333: 'QMS', 334: 'SIGN', 336: 'XRAYE'}
for r in range(80, 92):
  if r != 85: M[r] = 'LS-67'
for r in list(range(220, 252)): M[r] = 'LS-76'
EXTRA = {17: ['LS-42'], 138: ['LS-61'], 295: ['LS-84'], 296: ['LS-84']}
HOW = {'copy': 'بند المرجع كما هو', 'ls': 'بند المرجع معدّل للمشروع', 'new': 'بند جديد من المرجع الفني', 'opt': 'اختياري (كمية صفر)'}
BY = {x['k']: x for x in L}

NAVY, SKY, GOLD = '153E5C', 'E8F6FC', 'FFF3DF'
F = lambda **k: Font(name='Tajawal', **k)
thin = Side(style='thin', color='D5DEE6'); BR = Border(left=thin, right=thin, top=thin, bottom=thin)
def head(ws, row, cols):
  for j, (t, w) in enumerate(cols, 1):
    c = ws.cell(row, j, t); c.font = F(bold=True, color='FFFFFF', size=10.5); c.fill = PatternFill('solid', fgColor=NAVY); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BR
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
def cell(ws, r, j, v, **kw):
  c = ws.cell(r, j, v); c.font = F(size=10, **{k: v for k, v in kw.items() if k in ('bold', 'color')}); c.alignment = Alignment(vertical='top', wrap_text=True, horizontal=kw.get('h', 'right')); c.border = BR
  if kw.get('fill'): c.fill = PatternFill('solid', fgColor=kw['fill'])
  return c
def title(ws, t, sub, ncol):
  ws.sheet_view.rightToLeft = True
  ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol); c = ws.cell(1, 1, t); c.font = F(bold=True, size=15, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY); c.alignment = Alignment(horizontal='right', vertical='center'); ws.row_dimensions[1].height = 30
  ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol); c = ws.cell(2, 1, sub); c.font = F(size=10, color='5B6B7A'); c.alignment = Alignment(horizontal='right', wrap_text=True); ws.row_dimensions[2].height = 32
  ws.oddFooter.left.text = '© 2026 حقوق النشر محفوظة لتجمع الأحساء الصحي — إدارة الخدمات الفنية'; ws.oddFooter.left.size = 8; ws.oddFooter.right.text = '&P / &N'
  ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True

src = openpyxl.load_workbook(SOW, data_only=True)['SOW']
rows = []
for i, r in enumerate(src.iter_rows(values_only=True, max_col=7), 1):
  if r[2] in ('Item', 'NO') or (r[5] and str(r[5]).lower().strip() == 'optional'):
    no = r[0]; no = (f'{no:.2f}'.rstrip('0').rstrip('.') if isinstance(no, float) else str(no or '—'))
    rows.append((i, no, str(r[1]).strip(), str(r[5]).strip()))
missing = [i for i, *_ in rows if i not in M]
assert not missing, missing

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'المقابلة'
title(ws, 'مقابلة نطاق العمل ببنود جدول الكميات — مركز فاطمة العفالق', 'كل بند في نطاق العمل (SOW 25-10-2026) وأين يُغطّى في جدول الكميات المبني بطريقة التجمع (مقطوعية بنصوص المرجع الفني). رقم البند في العمود الخامس هو رقمه في جدول الكميات على المنصة.', 9)
cols = [('#', 5), ('بند نطاق العمل', 10), ('الوصف في نطاق العمل', 42), ('الحالة', 10), ('رقم بندنا', 8), ('بند جدول الكميات', 34), ('مرجع المقطوعية', 11), ('طريقة التغطية', 20), ('ملاحظة البند في الجدول', 60)]
head(ws, 4, cols); r = 5
ST = {'Included': 'مشمول', 'Excluded': 'مستبعد', 'optional': 'اختياري', 'Optional': 'اختياري'}
FILL = {'copy': None, 'ls': SKY, 'new': 'E7F5F0', 'opt': GOLD}
for n, (row, no, d, stt) in enumerate(rows, 1):
  k = M[row]; x = BY.get(k) if k else None
  vals = [n, no, d, ST.get(stt, stt), KEY[k] if k else '—', (x.get('title') or (k + ' — نص بند المقطوعية')) if x else 'غير مشمول في الجدول', x['ls'] if x and x['ls'] else '—', HOW[x['mode']] if x else 'مستبعد حسب نطاق العمل', (x['note'] + (' — ويغطيه أيضاً البند ' + '، '.join(str(KEY[e]) for e in EXTRA[row]) if row in EXTRA else '')) if x else 'مذكور في الملاحظات العامة للجدول.']
  for j, v in enumerate(vals, 1): cell(ws, r, j, v, fill=(FILL[x['mode']] if x else 'F3F5F8') if j >= 5 else None, h='center' if j in (1, 2, 4, 5, 7) else 'right')
  r += 1
ws.freeze_panes = 'A5'

# أسماء بنود النسخ من الإنتاج غير متاحة هنا؛ نأخذها من القائمة المحلية للعرض فقط
try: LSN = {c: t for c, t, _ in json.load(open('/tmp/ls_local.json'))}
except Exception: LSN = {}
for row in ws.iter_rows(min_row=5, min_col=6, max_col=6):
  for c in row:
    if str(c.value).endswith('نص بند المقطوعية'): code = str(c.value).split(' — ')[0]; c.value = LSN.get(code, code)

ws2 = wb.create_sheet('بنود الجدول')
title(ws2, f'بنود جدول الكميات ({len(L)} بنداً) ومصدر كل بند', 'البنود بالترتيب نفسه على المنصة. النص الكامل لكل بند ومواصفاته في جدول الكميات على المنصة، ويُصدَّر منها Excel أو دليل المواصفات PDF.', 7)
head(ws2, 4, [('#', 5), ('بند الجدول', 38), ('مرجع المقطوعية', 12), ('المصدر', 22), ('الكمية', 8), ('بنود نطاق العمل المغطاة', 26), ('ملاحظة البند', 64)])
cov = {}
for row, no, d, stt in rows:
  for k_ in ([M[row]] if M[row] else []) + EXTRA.get(row, []): cov.setdefault(k_, []).append(no)
r = 5
for i, x in enumerate(L, 1):
  t = x.get('title') or LSN.get(x['ls'], x['ls'])
  vals = [i, t, x['ls'] or '—', HOW[x['mode']], 0 if x['mode'] == 'opt' else 1, '، '.join(cov.get(x['k'], [])) or 'إضافة من مرجعنا', x['note']]
  for j, v in enumerate(vals, 1): cell(ws2, r, j, v, fill=FILL[x['mode']] if j == 4 else ('FFF3DF' if j == 6 and not cov.get(x['k']) else None), h='center' if j in (1, 3, 5) else 'right')
  r += 1
ws2.freeze_panes = 'A5'

ws3 = wb.create_sheet('للقرار')
title(ws3, 'فروقات وإضافات تحتاج قرارك', 'ما اختلف فيه نطاق العمل عن مرجعنا، وما أضفناه من مرجعنا ولم يذكره نطاق العمل.', 4)
head(ws3, 4, [('#', 5), ('الموضوع', 30), ('نطاق العمل', 45), ('ما في جدولنا والتوصية', 70)])
D = [
  ('سماكة الإنترلوك للممرات', 'Interlock 200×100×60 مم (2.2.3)', 'المرجع 80 مم للممرات و100 مم تحت المواقف (LS-89)؛ الجدول يعتمد الأعلى. القرار: إبقاء 80 مم أو النزول إلى 60 مم.'),
  ('نظام البلاطات', 'Solid slab (3.2.6)', 'بند الهيكل (LS-47) نصه للبلاطات اللاكمرية؛ أضفنا ملاحظة أن النظام حسب المخططات (مصمتة). يكفي إن كانت المخططات واضحة.'),
  ('نظام التكييف', 'وحدات هواء نقي DX + وحدات FCU DX + CAV/VAV (23.x)', 'بُني نص جديد على أساس LS-77 (المجاري والمخمدات والفلاتر وASHRAE 170) بدل وحدات Package وحدها، وأُضيفت أجهزة فرق الضغط لغرف العزل والإنعاش.'),
  ('السخانات', 'سخانات كهربائية 30 لتراً (22.1.5)', 'المرجع لسخانات مركزية 80 جالون؛ الجدول التزم بنطاق العمل (30 لتراً لنقاط الاستخدام) مع إبقاء A.O. Smith.'),
  ('الغازات الطبية', 'منظومة كاملة: O2 وN2O وMA4 وشفط وAGSS ومحطات ومشعبات (22.3)', 'المرجع كان للأكسجين فقط؛ كُتب البند كاملاً من نطاق العمل ومراجع المرجع الفني (22 6x) واختبارات HTM 02-01 / NFPA 99.'),
  ('المصاعد', 'مصعدان: الموظفين GF-16 والزوار GF-76 (14.x)', 'Schindler حصراً كما في سجل المواد؛ الحمولة والكابينة حسب أبعاد البئر. أكّد هل مصعد الموظفين للنقالات.'),
  ('وحدة الحلقة والمحول', 'RMU (Lucy) ومحول وكابلات جهد متوسط (26.1–26.2، 26.10)', 'بند جديد؛ Lucy كما ورد أو ما تعتمده الشركة السعودية للكهرباء.'),
  ('استدعاء التمريض', 'Nurse Call (27.6.1)', 'المرجع SOLT اللاسلكي للمراكز الصغيرة. إن كان المركز طوارئ عاجلة (UCC) بغرف ملاحظة، قرّر هل يكفي أو يُطلب نظام سلكي.'),
  ('إضافة من مرجعنا', '—', 'الدراسات والاختبارات المعملية (LS-43)، مستلزمات مكافحة العدوى الجدارية (LS-64)، التنظيف النهائي والوثائق والتدريب (LS-32)، توصيل الخدمات العامة (LS-92): ليست في نطاق العمل لكنها من اشتراطاتنا. احذف ما لا تريده.'),
  ('المستبعد', 'التجهيزات الطبية، المعدات الكهربائية', 'غير مشمولة في الجدول، ومذكورة في الملاحظات العامة. الأثاث غير الطبي كان مستبعداً وأُدرج بقرارك (بند الأثاث LS-93).'),
  ('الاختيارية', 'QMS، الشاشات الرقمية، تجهيز الأشعة الكهربائي', 'في الجدول بكمية صفر: يسعّرها المقاول ولا تدخل المجموع.')]
for i, (a, b, c) in enumerate(D, 1):
  for j, v in enumerate([i, a, b, c], 1): cell(ws3, 4 + i, j, v, bold=(j == 2), h='center' if j == 1 else 'right')

ws0 = wb.create_sheet('الملخص', 0)
title(ws0, 'جدول كميات مركز فاطمة العفالق — ملخص المقابلة', 'مبني على نطاق العمل (SOW) الصادر 25 أكتوبر 2026، بطريقة التجمع: بنود مقطوعية بنصوص المرجع الفني ومراجع MasterFormat.', 3)
head(ws0, 4, [('البيان', 46), ('العدد', 10), ('', 40)])
cnt = lambda m: sum(1 for x in L if x['mode'] == m)
S = [('بنود نطاق العمل', len(rows)), ('— مشمولة ومقابلة ببنود الجدول', sum(1 for r_, *_ in rows if M[r_] and BY[M[r_]]['mode'] != 'opt')), ('— اختيارية', sum(1 for r_, *_ in rows if M[r_] and BY[M[r_]]['mode'] == 'opt')), ('— مستبعدة', sum(1 for r_, *_ in rows if not M[r_])),
     ('بنود جدول الكميات', len(L)), ('— بند المرجع كما هو (بملاحظة للمشروع)', cnt('copy')), ('— بند المرجع معدّل للمشروع', cnt('ls')), ('— بند جديد من المرجع الفني', cnt('new')), ('— اختياري (كمية صفر)', cnt('opt')),
     ('بنود من مرجعنا لم يذكرها نطاق العمل', sum(1 for x in L if not cov.get(x['k'])))]
for i, (a, b) in enumerate(S, 5):
  cell(ws0, i, 1, a, bold=not a.startswith('—')); cell(ws0, i, 2, b, h='center', bold=True); cell(ws0, i, 3, '')
wb.save(OUT); print('ok', len(rows))
